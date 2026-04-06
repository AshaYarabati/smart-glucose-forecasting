# =============================================================================
# Project: Smart Glucose Forecasting (MSc Thesis)
# Author: Asha Deepthi Yarabati
# Description: Script to export merged results table to Word document for 
#              posting into thesis report. 
# Version: 1.0
# Last Updated: 06 Apr 2026
# Usage:
# - Run in project root directory "pwsh ./run_all.ps1"
# =============================================================================


from __future__ import annotations

import argparse
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from common import ensure_dir


def style_doc(doc: Document, title: str):
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(title)
    run.bold = True
    run.font.size = Pt(15)
    p2 = doc.add_paragraph('Consolidated test-set results across baseline and deep learning models.')
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph('')


def add_table(doc: Document, df: pd.DataFrame):
    table = doc.add_table(rows=1, cols=len(df.columns))
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table.rows[0].cells
    for i, col in enumerate(df.columns):
        p = hdr[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(str(col))
        r.bold = True
        r.font.size = Pt(10)
    for _, row in df.iterrows():
        cells = table.add_row().cells
        for i, val in enumerate(row):
            p = cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i != 0 else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(str(val))
            r.font.size = Pt(9)


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description='Export merged results table to DOCX.')
    parser.add_argument('--input', required=True)
    parser.add_argument('--output-dir', required=True)
    parser.add_argument('--output-name', default='merged_results_table.docx')
    parser.add_argument('--title', default='Smart Glucose Forecasting Results')
    args = parser.parse_args()

    out_dir = ensure_dir(args.output_dir)
    df = pd.read_csv(args.input)
    doc = Document()
    style_doc(doc, args.title)
    add_table(doc, df)
    out_path = Path(out_dir) / args.output_name
    doc.save(out_path)
    print(f'Wrote Word document: {out_path}')
